#!/usr/bin/env python3
"""Build the FORGE Nature Biotechnology pre-results working manuscript."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(23, 63, 115)
TEAL = RGBColor(75, 155, 150)
CORAL = RGBColor(212, 91, 101)
CHARCOAL = RGBColor(41, 47, 54)
GRAY = RGBColor(105, 112, 120)
PLACEHOLDER = RGBColor(127, 96, 0)


def set_run_font(
    run,
    *,
    name: str = "Arial",
    size: float = 12,
    bold: bool | None = None,
    italic: bool | None = None,
    color: RGBColor | None = None,
) -> None:
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_style_font(
    style,
    *,
    name: str,
    size: float,
    bold: bool = False,
    italic: bool = False,
    color: RGBColor = CHARCOAL,
) -> None:
    style.font.name = name
    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = color


def add_keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepNext")) is None:
        p_pr.append(OxmlElement("w:keepNext"))


def add_keep_together(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepLines")) is None:
        p_pr.append(OxmlElement("w:keepLines"))


def add_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_paragraph_border(paragraph, color: str = "B8CDE3") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    for edge_name in ("top", "bottom", "left", "right"):
        edge = OxmlElement(f"w:{edge_name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:space"), "5")
        edge.set(qn("w:color"), color)
        borders.append(edge)


def highlight_run(run, fill: str = "FFF2CC") -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    run._element.get_or_add_rPr().append(shd)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=9, color=GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def add_continuous_line_numbers(section) -> None:
    sect_pr = section._sectPr
    for existing in sect_pr.findall(qn("w:lnNumType")):
        sect_pr.remove(existing)
    line_numbers = OxmlElement("w:lnNumType")
    line_numbers.set(qn("w:countBy"), "1")
    line_numbers.set(qn("w:distance"), "360")
    line_numbers.set(qn("w:restart"), "continuous")
    columns = sect_pr.find(qn("w:cols"))
    if columns is None:
        sect_pr.append(line_numbers)
    else:
        sect_pr.insert(sect_pr.index(columns), line_numbers)


def enable_field_updates(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    set_style_font(normal, name="Arial", size=12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    title = document.styles["Title"]
    set_style_font(title, name="Arial", size=18, bold=True, color=NAVY)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.line_spacing = 1.15

    h1 = document.styles["Heading 1"]
    set_style_font(h1, name="Arial", size=13, bold=True, color=NAVY)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(2)
    h1.paragraph_format.line_spacing = 1.15

    h2 = document.styles["Heading 2"]
    set_style_font(h2, name="Arial", size=11.5, bold=True, color=NAVY)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(1)
    h2.paragraph_format.line_spacing = 1.15

    if "Manuscript note" not in document.styles:
        note = document.styles.add_style("Manuscript note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        note = document.styles["Manuscript note"]
    set_style_font(note, name="Arial", size=9.5, color=GRAY)
    note.paragraph_format.space_before = Pt(5)
    note.paragraph_format.space_after = Pt(5)
    note.paragraph_format.left_indent = Inches(0.12)
    note.paragraph_format.right_indent = Inches(0.12)
    note.paragraph_format.line_spacing = 1.2


def configure_document(document: Document) -> None:
    document.core_properties.title = (
        "Generative design of synthesizable ionizable lipids for mRNA delivery"
    )
    document.core_properties.subject = "Nature Biotechnology pre-results working manuscript"
    document.core_properties.keywords = "FORGE, ionizable lipids, LNP, mRNA, generative model"
    enable_field_updates(document)

    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.08)
        section.right_margin = Inches(0.92)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)
        add_continuous_line_numbers(section)

        header = section.header
        header.is_linked_to_previous = True
        header_p = header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_run = header_p.add_run("FORGE  |  pre-results working manuscript")
        set_run_font(header_run, name="Arial", size=8.5, color=GRAY)

        footer = section.footer
        footer.is_linked_to_previous = True
        add_page_number(footer.paragraphs[0])


INLINE_TOKEN = re.compile(r"(\*\*.*?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+\])")
CITATION = re.compile(r"^\[(\d+(?:[-,]\d+)*)\]$")


def add_inline(paragraph, text: str, *, default_size: float = 12) -> None:
    for token in INLINE_TOKEN.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=default_size, bold=True)
            continue
        if token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=default_size, italic=True)
            continue
        if token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Courier New", size=max(default_size - 1, 8.5))
            continue
        citation = CITATION.match(token)
        if citation:
            run = paragraph.add_run(citation.group(1))
            set_run_font(run, size=max(default_size - 2.5, 8))
            run.font.superscript = True
            continue
        if token.startswith("[") and token.endswith("]"):
            run = paragraph.add_run(token)
            set_run_font(run, size=default_size, bold=True, color=PLACEHOLDER)
            highlight_run(run)
            continue
        run = paragraph.add_run(token)
        set_run_font(run, size=default_size)


def add_callout(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Manuscript note")
    add_paragraph_shading(paragraph, "EEF5FB" if "FIGURE" in text else "FFF8E8")
    add_paragraph_border(paragraph, "AFC9E2" if "FIGURE" in text else "E5C56B")
    add_keep_together(paragraph)
    add_inline(paragraph, text, default_size=9.5)


def add_body_paragraph(document: Document, text: str, *, reference: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    paragraph.paragraph_format.space_after = Pt(0)
    if reference:
        paragraph.paragraph_format.left_indent = Inches(0.26)
        paragraph.paragraph_format.first_line_indent = Inches(-0.26)
    add_inline(paragraph, text)


def normalize_block(lines: list[str]) -> str:
    return " ".join(line.strip() for line in lines).strip()


def count_words(text: str) -> int:
    text = re.sub(r"<!--.*?-->", " ", text, flags=re.S)
    text = re.sub(r"[#>*_`]", " ", text)
    text = re.sub(r"\[[^\]]*PLACEHOLDER[^\]]*\]", " ", text)
    return len(re.findall(r"\b[\w'-]+\b", text))


def split_once(source: str, *markers: str) -> tuple[str, str] | None:
    """Split on the first marker present, so heading level changes do not break the build."""

    for marker in markers:
        if marker in source:
            head, tail = source.split(marker, 1)
            return head, tail
    return None


def manuscript_counts(source: str) -> tuple[int, int]:
    after_abstract = split_once(source, "# Abstract\n", "## Abstract\n")
    if after_abstract is None:
        raise SystemExit("no Abstract heading found in the manuscript")
    intro = split_once(after_abstract[1], "# Introduction",
                       "<!-- Unheaded introduction begins here. -->")
    if intro is None:
        raise SystemExit("no Introduction marker found in the manuscript")
    abstract, rest = intro
    main = rest.split("# Online Methods", 1)[0]
    # The supplement is generated and sits after References; it is not main text.
    main = main.split("<!-- SUPPLEMENT:BEGIN", 1)[0]
    return count_words(abstract), count_words(main)


def build_document(source_path: Path, output_path: Path) -> None:
    source = source_path.read_text(encoding="utf-8")
    abstract_words, main_words = manuscript_counts(source)
    document = Document()
    configure_styles(document)
    configure_document(document)

    lines = source.splitlines()
    index = 0
    title_seen = False
    in_references = False
    before_abstract = True

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if not stripped or stripped.startswith("<!--"):
            index += 1
            continue

        if stripped.startswith(">"):
            block: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                block.append(lines[index].strip()[1:].strip())
                index += 1
            add_callout(document, normalize_block(block))
            continue

        if stripped.startswith("# "):
            heading = stripped[2:].strip()
            if not title_seen:
                title_p = document.add_paragraph(style="Title")
                add_inline(title_p, heading, default_size=18)
                title_seen = True
            else:
                in_references = heading == "References"
                heading_p = document.add_paragraph(heading, style="Heading 1")
                if heading in {"Online Methods", "Figure legends", "References"}:
                    heading_p.paragraph_format.page_break_before = True
                add_keep_with_next(heading_p)
            index += 1
            continue

        if stripped.startswith("## "):
            heading = stripped[3:].strip()
            if heading == "Abstract":
                before_abstract = False
                heading_p = document.add_paragraph("Abstract", style="Heading 1")
                heading_p.paragraph_format.space_before = Pt(8)
                add_keep_with_next(heading_p)
                meta = document.add_paragraph(style="Manuscript note")
                meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                meta_run = meta.add_run(
                    f"Abstract: {abstract_words} words  |  Main text: {main_words} words"
                )
                set_run_font(meta_run, name="Arial", size=8.5, color=GRAY)
            else:
                heading_p = document.add_paragraph(heading, style="Heading 2")
                add_keep_with_next(heading_p)
            index += 1
            continue

        block = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate.startswith("#")
                or candidate.startswith(">")
                or candidate.startswith("<!--")
            ):
                break
            block.append(candidate)
            index += 1
        paragraph_text = normalize_block(block)

        if before_abstract and paragraph_text.startswith("**[Author"):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.line_spacing = 1.2
            add_inline(paragraph, paragraph_text)
            continue
        if before_abstract and paragraph_text.startswith("*[Affiliations"):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.line_spacing = 1.2
            add_inline(paragraph, paragraph_text)
            continue

        is_reference = in_references and re.match(r"^\d+\.\s", paragraph_text) is not None
        add_body_paragraph(document, paragraph_text, reference=is_reference)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source",
        type=Path,
        default=Path("manuscript/FORGE_Nature_Biotechnology_working_draft.md"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("manuscript/FORGE_Nature_Biotechnology_working_draft.docx"),
    )
    args = parser.parse_args()
    build_document(args.source, args.output)
    print(args.output)


if __name__ == "__main__":
    main()
